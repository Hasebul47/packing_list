"""
unified_packing.py
------------------
GAP / Old Navy PO → Packing List (Unified)

Logic  : GAP (Bulk / Single / Multi Y / Multi N) — merge same STYLE by COLOR
UI     : Interactive preview (yellow strips, PO header, live totals, Excel)

Run:
    streamlit run unified_packing.py
"""

from __future__ import annotations

import io
import json
from pathlib import Path

import pandas as pd
import streamlit as st

import parser_lib as parser
from color_builders import (
    BULK_CTN_CAP,
    DEFAULT_CTN_CAP,
    DEFAULT_CTN_MEAS,
    DEFAULT_EXCESS_PCT,
    DEFAULT_NW,
    build_bulk_by_color,
    build_multi_y_by_color,
    build_single_by_color,
    detect_table_type,
    section_to_display_rows,
    sort_sizes,
)
from multi_n_core import (
    DEFAULT_CTN_CODE,
    DEFAULT_NET_DEDUCT,
    DEFAULT_NNW,
    DEFAULT_PREPACKS_PER_CTN,
    DEFAULT_TARE,
    build_multi_n,
)
from multi_n_export import export_multi_n_excel

PACK_TYPES = ("Bulk", "Single", "Multi Y", "Multi N")


def detect_all(data: dict) -> dict[str, list[str]]:
    found = {t: [] for t in PACK_TYPES}
    for tname, table in (data.get("tables") or {}).items():
        ptype = detect_table_type(table)
        if ptype:
            found[ptype].append(tname)
    return found


def inject_css():
    st.markdown(
        """
<style>
  .plist-wrap { font-family: 'Courier New', monospace; font-size: 13px; }
  .cc-strip {
    text-align: center; font-weight: bold; font-size: 15px;
    letter-spacing: 2px; background: #FFFF00; border: 1px solid #000;
    padding: 6px 0; margin: 14px 0 8px 0;
  }
  .po-grid {
    display: grid; grid-template-columns: 1.3fr 1fr; gap: 18px;
    margin-bottom: 8px; line-height: 1.45;
  }
  .po-line { display: flex; }
  .po-label { width: 110px; font-weight: bold; }
  .po-colon { width: 14px; text-align: center; }
  table.plist {
    border-collapse: collapse; width: 100%; margin: 6px 0 14px 0;
    font-family: 'Courier New', monospace; font-size: 12.5px;
  }
  table.plist th, table.plist td {
    border: 1px solid #333; padding: 3px 6px; text-align: center;
  }
  table.plist th { background: #4472C4; color: #fff; font-weight: bold; }
  table.plist tr.total td { background: #FFF2CC; font-weight: bold; }
</style>
""",
        unsafe_allow_html=True,
    )


def render_section_strip(title: str):
    st.markdown(f'<div class="cc-strip">{title}</div>', unsafe_allow_html=True)


def render_po_header_html(sec: dict, buyer: str, po_no: str, grand_ctn: int):
    meas = "".join(
        f'<div class="po-line"><span class="po-label">CTN MEAS.</span>'
        f'<span class="po-colon">:</span><span>{m}</span></div>'
        for m in (sec.get("ctn_meas_list") or DEFAULT_CTN_MEAS)
    )
    ctn_from = sec["rows"][0]["ctn_from"] if sec.get("rows") else 1
    exs = sec["ship_qty"] - sec["qty_ordered"]
    html = f"""
<div class="plist-wrap">
  <div class="po-grid">
    <div>
      <div class="po-line"><span class="po-label">BUYER</span><span class="po-colon">:</span><span>{buyer}</span></div>
      <div class="po-line"><span class="po-label">STYLE</span><span class="po-colon">:</span><span>{sec.get('style_full','')}</span></div>
      <div class="po-line"><span class="po-label">P.O. #</span><span class="po-colon">:</span><span>{po_no}</span></div>
      <div class="po-line"><span class="po-label">O/QTY</span><span class="po-colon">:</span><span>{sec['qty_ordered']} PCS</span></div>
      <div class="po-line"><span class="po-label">SHIP QTY</span><span class="po-colon">:</span><span>{sec['ship_qty']} PCS</span></div>
      <div class="po-line"><span class="po-label">EX/SHORT</span><span class="po-colon">:</span><span>{exs} PCS</span></div>
      <div class="po-line"><span class="po-label">CTN QTY</span><span class="po-colon">:</span><span>{sec['ctn_total']} CTNS</span></div>
      {meas}
    </div>
    <div style="margin-top:48px">
      <div class="po-line"><span class="po-label">PO #</span><span class="po-colon">:</span><span>{po_no}</span></div>
      <div class="po-line"><span class="po-label">SKU / Item</span><span class="po-colon">:</span><span>{sec.get('style_full','')}</span></div>
      <div class="po-line"><span class="po-label">Unit / Prepack</span><span class="po-colon">:</span><span>{sec.get('units_per_prepack',1)}</span></div>
      <div class="po-line"><span class="po-label">CARTON</span><span class="po-colon">:</span><span>{ctn_from:02d} of {grand_ctn}</span></div>
    </div>
  </div>
</div>
"""
    st.markdown(html, unsafe_allow_html=True)


def render_plist_table(sec: dict):
    recs = section_to_display_rows(sec)
    if not recs:
        st.warning("No rows")
        return
    cols = list(recs[0].keys())
    th = "".join(f"<th>{c}</th>" for c in cols)
    body = ""
    for rec in recs:
        cls = ' class="total"' if str(rec.get("CTN NO", "")).upper() == "TOTAL" else ""
        tds = "".join(f"<td>{rec.get(c, '')}</td>" for c in cols)
        body += f"<tr{cls}>{tds}</tr>"
    st.markdown(
        f'<div class="plist-wrap"><table class="plist"><tr>{th}</tr>{body}</table></div>',
        unsafe_allow_html=True,
    )


def export_unified_excel(data, bulk_secs, single_secs, multi_y_secs, multi_n_secs) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    wb = Workbook()
    thin = Border(left=Side(style="thin"), right=Side(style="thin"),
                  top=Side(style="thin"), bottom=Side(style="thin"))
    header_fill = PatternFill("solid", fgColor="4472C4")
    header_font = Font(bold=True, color="FFFFFF", name="Courier New", size=10)
    total_fill = PatternFill("solid", fgColor="FFF2CC")
    yellow = PatternFill("solid", fgColor="FFFF00")
    type_fills = {
        "Bulk": PatternFill("solid", fgColor="D9EAD3"),
        "Single": PatternFill("solid", fgColor="FCE4D6"),
        "Multi Y": PatternFill("solid", fgColor="D0E0F0"),
        "Multi N": PatternFill("solid", fgColor="EAD1DC"),
    }
    font_bold = Font(bold=True, name="Courier New", size=11)
    font_norm = Font(name="Courier New", size=10)
    font_title = Font(bold=True, name="Courier New", size=14)
    ctr = Alignment(horizontal="center", vertical="center")

    header = data.get("PO_HEADER") or {}
    buyer = header.get("BUYER") or "OLD NAVY"
    style = header.get("STYLE") or ""
    po_no = str(header.get("P.O._#") or header.get("P.O. #") or "")
    oiqty = header.get("OIQTY") or 0
    all_secs = bulk_secs + single_secs + multi_y_secs + multi_n_secs
    grand_ctn = sum(s.get("ctn_total", 0) for s in all_secs)

    ws = wb.active
    ws.title = "Summary"
    ws["A1"] = "CREATIVE COLLECTIONS LTD"
    ws["A1"].font = font_title
    ws["A2"] = "PACKING LIST — UNIFIED"
    ws["A3"] = f"BUYER: {buyer}"
    ws["A4"] = f"STYLE: {style}"
    ws["A5"] = f"PO #: {po_no}"
    ws["A6"] = f"O/QTY: {oiqty} PCS"
    ws["A7"] = f"TOTAL CTNS: {grand_ctn}"
    ws["A9"] = "Detected packing types"
    ws["A9"].font = font_bold
    row = 10
    for label, secs in [("Bulk", bulk_secs), ("Single", single_secs),
                        ("Multi Y", multi_y_secs), ("Multi N", multi_n_secs)]:
        n_ctn = sum(s.get("ctn_total", 0) for s in secs)
        n_qty = sum(s.get("ship_qty", 0) for s in secs)
        ws.cell(row, 1, label)
        ws.cell(row, 2, f"{len(secs)} color(s)")
        ws.cell(row, 3, f"{n_ctn} CTNS")
        ws.cell(row, 4, f"{n_qty} PCS")
        if label in type_fills:
            for c in range(1, 5):
                ws.cell(row, c).fill = type_fills[label]
        row += 1

    def write_sections_sheet(sheet_name, sections):
        if not sections:
            return
        ws = wb.create_sheet(sheet_name[:31])
        r = 1
        ws.cell(r, 1, "CREATIVE COLLECTIONS LTD-1A.").font = font_title
        r = 2
        ws.cell(r, 1, "Nishat Nagar , Tongi , Gazipur .").font = font_norm
        r = 4
        for sec in sections:
            ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=18)
            cell = ws.cell(r, 1, sec["pack_type"].upper().replace("MULTI Y", "MULTY PACK (Y)"))
            cell.fill = yellow
            cell.font = Font(bold=True, name="Courier New", size=12)
            cell.alignment = ctr
            r += 1
            left_lines = [
                ("BUYER", buyer), ("STYLE", sec.get("style_full") or style),
                ("P.O. #", po_no), ("O/QTY", f"{sec['qty_ordered']} PCS"),
                ("SHIP QTY", f"{sec['ship_qty']} PCS"),
                ("EX/SHORT", f"{sec['ship_qty'] - sec['qty_ordered']} PCS"),
                ("CTN QTY", f"{sec['ctn_total']} CTNS"),
            ]
            for label, val in left_lines:
                ws.cell(r, 1, label).font = font_bold
                ws.cell(r, 2, ":")
                ws.cell(r, 3, val).font = font_norm
                r += 1
            for m in sec.get("ctn_meas_list") or DEFAULT_CTN_MEAS:
                ws.cell(r, 1, "CTN MEAS.").font = font_bold
                ws.cell(r, 2, ":")
                ws.cell(r, 3, m).font = font_norm
                r += 1
            right_r = r - len(left_lines) - len(sec.get("ctn_meas_list") or DEFAULT_CTN_MEAS)
            ws.cell(right_r, 10, "PO #").font = font_bold
            ws.cell(right_r, 11, po_no)
            ws.cell(right_r + 1, 10, "SKU / Item").font = font_bold
            ws.cell(right_r + 1, 11, sec.get("style_full") or style)
            ws.cell(right_r + 2, 10, "Unit / Prepack").font = font_bold
            ws.cell(right_r + 2, 11, str(sec.get("units_per_prepack", 1)))
            ws.cell(right_r + 3, 10, "CARTON").font = font_bold
            ctn_from = sec["rows"][0]["ctn_from"] if sec["rows"] else 1
            ws.cell(right_r + 3, 11, f"{ctn_from:02d} of {grand_ctn}")
            r += 1
            recs = section_to_display_rows(sec)
            if not recs:
                r += 2
                continue
            cols = list(recs[0].keys())
            for c, col in enumerate(cols, 1):
                cell = ws.cell(r, c, col)
                cell.font = header_font
                cell.fill = header_fill
                cell.border = thin
                cell.alignment = ctr
            r += 1
            for rec in recs:
                is_total = str(rec.get("CTN NO", "")).upper() == "TOTAL"
                for c, col in enumerate(cols, 1):
                    cell = ws.cell(r, c, rec.get(col, ""))
                    cell.border = thin
                    cell.alignment = ctr
                    cell.font = font_bold if is_total else font_norm
                    if is_total:
                        cell.fill = total_fill
                r += 1
            r += 2
        for col in ws.columns:
            letter = col[0].column_letter
            maxlen = max((len(str(cell.value)) for cell in col if cell.value), default=0)
            ws.column_dimensions[letter].width = min(maxlen + 2, 22)

    write_sections_sheet("Bulk", bulk_secs)
    write_sections_sheet("Single", single_secs)
    write_sections_sheet("Multi_Y", multi_y_secs)
    write_sections_sheet("Multi_N", multi_n_secs)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ===========================================================================
# Streamlit UI
# ===========================================================================

def run(pdf_file=None):
    """Render the GAP / Old Navy page. pdf_file: uploaded PDF (or None)."""
    inject_css()
    st.title("📦 GAP / Old Navy — PO → Packing List")
    st.caption(
        "Upload PDF → extract → auto-detect **Bulk / Single / Multi Y / Multi N** → "
        "merge same STYLE by **COLOR** → interactive packing list + Excel"
    )

    uploaded = pdf_file
    with st.expander("Or load an existing JSON"):
        json_file = st.file_uploader("JSON file", type=["json"], key="json_up")

    data = None
    source_name = ""

    if uploaded is not None:
        source_name = uploaded.name
        sig = f"{uploaded.name}::{uploaded.size}"
        if st.session_state.get("gap_upload_sig") != sig:
            for k in list(st.session_state.keys()):
                if k.startswith(("gap_xlsx_", "mn_xlsx_")):
                    st.session_state.pop(k, None)
            st.session_state["gap_upload_sig"] = sig
        with st.spinner("Reading PDF…"):
            try:
                text = parser.extract_text(io.BytesIO(uploaded.getvalue()))
            except Exception as e:
                st.error(f"PDF read failed: {e}")
                st.stop()
        with st.spinner("Parsing PO structure…"):
            try:
                data = parser.parse_text(text)
            except Exception as e:
                st.error(f"Parse failed: {e}")
                st.stop()
        st.success(
            f"Found **{len(data.get('tables', {}))}** table(s) "
            f"from `{uploaded.name}` ({len(text):,} characters)"
        )
    elif json_file is not None:
        source_name = json_file.name
        try:
            data = json.loads(json_file.getvalue().decode("utf-8"))
            st.success(f"Loaded JSON `{json_file.name}`")
        except Exception as e:
            st.error(f"Invalid JSON: {e}")
            st.stop()

    if data is None:
        st.info("Upload a GAP / Old Navy Destination PO PDF to get started.")
        st.markdown("""
### Logic (GAP)
| Type | Detection | Grouping |
|------|-----------|----------|
| **Bulk** | `PrePack Type == Bulk` | One section **per color** |
| **Single** | `PrePack Type == Single` | One section **per color** (prepacks merged) |
| **Multi Y** | `Multi` + `Full Carton == Y` | One section **per color** (all sizes one row) |
| **Multi N** | `Multi` + `Full Carton == N` | Per prepack (shared cartons) |

Same **STYLE** → tables merged by **COLOR** only.
    """)
        st.stop()

    header = data.get("PO_HEADER") or {}
    buyer = header.get("BUYER") or "OLD NAVY"
    style = header.get("STYLE") or ""
    po_no = str(header.get("P.O._#") or header.get("P.O. #") or "")
    oiqty = header.get("OIQTY") or 0

    detected = detect_all(data)
    active_types = [t for t in PACK_TYPES if detected[t]]

    m1, m2, m3, m4 = st.columns(4)
    m1.metric("BUYER", buyer)
    m2.metric("STYLE", style)
    m3.metric("PO #", po_no)
    m4.metric("O/QTY", f"{oiqty} PCS")

    st.markdown("### Auto-detected packing types")
    dcols = st.columns(4)
    for i, ptype in enumerate(PACK_TYPES):
        n = len(detected[ptype])
        if n:
            dcols[i].success(f"**{ptype}** — {n} table(s)")
        else:
            dcols[i].info(f"{ptype} — none")

    if not active_types:
        st.warning("No known packing types found.")
        st.json(data)
        st.stop()

    st.download_button(
        "⬇ Download extracted JSON",
        data=json.dumps(data, indent=2, ensure_ascii=False),
        file_name=f"{po_no or Path(source_name).stem}.json",
        mime="application/json",
    )

    with st.sidebar:
        st.header("Packing defaults")
        prepacks_per_ctn = st.number_input("Prepacks / carton (Multi N)", min_value=1, max_value=50, value=DEFAULT_PREPACKS_PER_CTN)
        single_max_pp = st.number_input(
            "Max prepacks / carton (Single) — 0 = auto by size", min_value=0, max_value=50, value=0)
        excess_pct = st.number_input(
            "Max excess % per size (Bulk)", min_value=0.0, max_value=10.0,
            value=float(DEFAULT_EXCESS_PCT), step=0.5,
            help="Extra pieces shipped with Bulk = floor(this % of each size's total order).")
        _sizes = sort_sizes(sorted({
            str(r.get("Size Desc") or "")
            for t in (data.get("tables") or {}).values() for r in (t.get("rows") or [])
        } - {""}))
        caps = {}
        with st.expander("Carton capacity (pcs per carton)"):
            for _s in _sizes:
                caps[_s] = st.number_input(
                    _s, min_value=1, max_value=500,
                    value=int(BULK_CTN_CAP.get(_s, DEFAULT_CTN_CAP)), key=f"cap_{_s}")
        net_deduct = st.number_input("Net deduct (Multi N)", value=float(DEFAULT_NET_DEDUCT), step=0.01)
        ctn_code = st.text_input("Carton code (Multi N)", value=DEFAULT_CTN_CODE)

    next_ctn = 1
    bulk_secs = build_bulk_by_color(data, start_ctn=next_ctn, caps=caps, excess_pct=excess_pct) if detected["Bulk"] else []
    next_ctn += sum(s["ctn_total"] for s in bulk_secs)
    single_secs = build_single_by_color(data, start_ctn=next_ctn, max_pp=int(single_max_pp) or None, caps=caps) if detected["Single"] else []
    next_ctn += sum(s["ctn_total"] for s in single_secs)
    multi_y_secs = build_multi_y_by_color(data, start_ctn=next_ctn) if detected["Multi Y"] else []
    next_ctn += sum(s["ctn_total"] for s in multi_y_secs)

    multi_n_secs = []
    multi_n_xlsx = None
    if detected["Multi N"]:
        try:
            multi_n_secs = build_multi_n(
                data, nw=DEFAULT_NW, tare_by_code=DEFAULT_TARE, net_deduct=net_deduct,
                default_per_ctn=prepacks_per_ctn, ctn_code=ctn_code, start_ctn_no=next_ctn,
            )
            grand = next_ctn - 1 + sum(s.get("ctn_total", 0) for s in multi_n_secs)
            multi_n_xlsx = export_multi_n_excel(
                multi_n_secs, DEFAULT_NW, DEFAULT_NNW, DEFAULT_TARE, net_deduct,
                list(DEFAULT_CTN_MEAS), buyer, po_no, grand,
            )
        except Exception as e:
            st.error(f"Multi N build failed: {e}")

    all_secs = bulk_secs + single_secs + multi_y_secs + multi_n_secs
    grand_ctn = sum(s.get("ctn_total", 0) for s in all_secs)
    grand_pcs = sum(s.get("ship_qty", 0) for s in all_secs)

    st.markdown("### Live totals")
    t1, t2, t3, t4 = st.columns(4)
    t1.metric("Total Cartons", grand_ctn)
    t2.metric("Total Pieces", grand_pcs)
    t3.metric("Order Qty", oiqty)
    exs = grand_pcs - oiqty
    if exs > 0:
        t4.success(f"Over-ship +{exs}")
    elif exs < 0:
        t4.error(f"Short-ship {exs}")
    else:
        t4.info("Exact match")

    tab_labels, tab_data = [], []
    if bulk_secs:
        tab_labels.append(f"Bulk ({len(bulk_secs)})")
        tab_data.append(("Bulk", bulk_secs))
    if single_secs:
        tab_labels.append(f"Single ({len(single_secs)})")
        tab_data.append(("Single", single_secs))
    if multi_y_secs:
        tab_labels.append(f"Multi Y ({len(multi_y_secs)})")
        tab_data.append(("Multi Y", multi_y_secs))
    if multi_n_secs:
        tab_labels.append(f"Multi N ({len(multi_n_secs)})")
        tab_data.append(("Multi N", multi_n_secs))
    tab_labels += ["Raw tables", "JSON"]
    tabs = st.tabs(tab_labels)

    for i, (ptype, secs) in enumerate(tab_data):
        with tabs[i]:
            strip_title = {
                "Bulk": "BULK PACK", "Single": "SINGLE PACK",
                "Multi Y": "MULTY PACK (Y)", "Multi N": "MULTY PACK (N)",
            }.get(ptype, ptype)
            for sec in secs:
                render_section_strip(strip_title)
                render_po_header_html(sec, buyer, po_no, grand_ctn)
                if ptype == "Multi N":
                    sizes = sec.get("sizes") or []
                    recs = []
                    for r in sec.get("rows") or []:
                        rec = {
                            "CTN NO": (f"{r['ctn_from']}-{r['ctn_to']}" if r.get("ctn_to") != r.get("ctn_from") else str(r.get("ctn_from"))),
                            "CTN MES": r.get("ctn_mes", ""), "CTN QTY": r.get("ctn_qty"),
                            "COLOR": sec.get("color"), "PREPACK": sec.get("prepack"),
                        }
                        for s in sizes:
                            rec[s] = (sec.get("units") or {}).get(s, "")
                        rec["PER BLST"] = r.get("per_blst", "")
                        rec["QTY PER CTN"] = r.get("qty_per_ctn", "")
                        rec["TOTAL QTY"] = r.get("total_qty", "")
                        rec["GROSS"] = r.get("gross", "")
                        rec["NET"] = r.get("net", "")
                        recs.append(rec)
                    st.dataframe(pd.DataFrame(recs), use_container_width=True, hide_index=True)
                else:
                    render_plist_table(sec)
                if sec.get("excess_by_size"):
                    st.caption("Excess shipped per size: " + ", ".join(
                        f"{k} +{v}" for k, v in sec["excess_by_size"].items()))
                c1, c2, c3 = st.columns(3)
                c1.metric("Ship qty", f"{sec['ship_qty']} PCS")
                c2.metric("Cartons", sec["ctn_total"])
                c3.metric("Color", sec.get("color", ""))
                st.divider()
            if ptype == "Multi N" and multi_n_xlsx:
                st.download_button(
                    "⬇ Download Multi N Excel (full template)", data=multi_n_xlsx,
                    file_name=f"packing_list_multi_n_{po_no}.xlsx",
                    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", key="dl_mn",
                )

    with tabs[len(tab_data)]:
        for tname, table in (data.get("tables") or {}).items():
            ptype = detect_table_type(table) or "?"
            rows = table.get("rows") or []
            with st.expander(f"{tname} — {ptype} — {len(rows)} rows"):
                if rows:
                    st.dataframe(rows, use_container_width=True)
                if table.get("total"):
                    st.json(table["total"])

    with tabs[len(tab_data) + 1]:
        st.json(data)

    st.markdown("---")
    st.markdown("### Download packing list")
    dl1, dl2 = st.columns(2)
    with dl1:
        if st.button("Prepare unified Excel", type="primary"):
            try:
                st.session_state["gap_xlsx_unified"] = export_unified_excel(
                    data, bulk_secs, single_secs, multi_y_secs, multi_n_secs
                )
                st.success("Excel ready")
            except Exception as e:
                st.error(f"Excel failed: {e}")
        if st.session_state.get("gap_xlsx_unified"):
            st.download_button(
                "⬇ Download unified packing list (.xlsx)",
                data=st.session_state["gap_xlsx_unified"],
                file_name=f"packing_list_unified_{po_no or 'PO'}.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                use_container_width=True,
            )
    with dl2:
        if multi_n_xlsx:
            st.download_button(
                "⬇ Multi N template Excel only", data=multi_n_xlsx,
                file_name=f"packing_list_multi_n_{po_no}.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                use_container_width=True, key="dl_mn2",
            )
        else:
            st.caption("Multi N template appears when Multi N tables are present.")


if __name__ == "__main__":
    st.set_page_config(page_title="GAP PO → Packing List", layout="wide", page_icon="📦")
    run(st.file_uploader("Upload Destination Purchase Order PDF", type=["pdf"]))
