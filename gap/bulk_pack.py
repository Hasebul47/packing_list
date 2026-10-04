# """
# bulk_pack.py
# ------------
# Streamlit app: read parsed PO JSON → build BULK PACK packing list.
# Each COLOR gets its own separate table.
# """

# import io
# import json
# import math
# from pathlib import Path

# import pandas as pd
# import streamlit as st


# st.set_page_config(page_title="Packing List — Bulk Pack", layout="wide")
# st.title("📦 Packing List — Bulk Pack")

# DEFAULT_JSON = "61480360.json"
# SIZE_ORDER = ["XS", "S", "M", "L", "XL", "XXL", "XXL+"]

# LEFT_STATIC  = ["CTN NO", "CTN MES:", "CTN QTY", "QTY", "COLOR", "PREPACK STECKER"]
# RIGHT_STATIC = ["QTY PER CTN", "TOTAL QTY", "GROSS WEIGHT", "NET WEIGHT"]


# @st.cache_data(show_spinner=False)
# def load_json(path: str) -> dict:
#     with open(path, "r", encoding="utf-8") as f:
#         return json.load(f)


# def color_display(raw: str) -> str:
#     if not raw:
#         return ""
#     parts = raw.rsplit(" ", 1)
#     return parts[-1] if len(parts) > 1 else raw


# def build_bulk_pack_grouped(data: dict) -> dict:
#     groups: dict[str, list] = {}
#     next_ctn: dict[str, int] = {}

#     for tname, table in data["tables"].items():
#         trows = table.get("rows") or []
#         if not trows:
#             continue
#         first = trows[0]
#         if first.get("PrePack Type") != "Bulk":
#             continue

#         color = color_display(first.get("Universal CC #Color Desc", "")) or "UNKNOWN"
#         if color not in next_ctn:
#             next_ctn[color] = 1

#         for row in trows:
#             size = row.get("Size Desc")
#             qty = row.get("Qty Ordered (eaches)") or 0
#             pack_factor = row.get("Item Carton Pack Factor")

#             cartons = math.ceil(qty / pack_factor) if pack_factor else None

#             start = next_ctn[color]
#             end = (start + cartons - 1) if cartons else start
#             next_ctn[color] = end + 1

#             data_row = {
#                 "CTN NO":          start,
#                 "CTN MES:":        end,
#                 "CTN QTY":         None,
#                 "QTY":             cartons,
#                 "COLOR":           color,
#                 "PREPACK STECKER": None,
#             }
#             for s in SIZE_ORDER:
#                 data_row[s] = qty if s == size else None
#             data_row["QTY PER CTN"]  = pack_factor
#             data_row["TOTAL QTY"]    = qty
#             data_row["GROSS WEIGHT"] = None
#             data_row["NET WEIGHT"]   = None

#             groups.setdefault(color, []).append(data_row)

#     # append TOTAL row per color
#     for color, rows in groups.items():
#         total_cartons = sum((r.get("QTY") or 0) for r in rows)
#         total_qty = sum((r.get("TOTAL QTY") or 0) for r in rows)

#         tot = {
#             "CTN NO":          "TOTAL",
#             "CTN MES:":        None,
#             "CTN QTY":         f"{total_cartons} CTN" if total_cartons else None,
#             "QTY":             None,
#             "COLOR":           None,
#             "PREPACK STECKER": None,
#         }
#         for s in SIZE_ORDER:
#             pieces = sum((r.get(s) or 0) for r in rows)
#             tot[s] = pieces if pieces else None
#         tot["QTY PER CTN"] = None
#         tot["TOTAL QTY"] = f"{total_qty} PCS" if total_qty else None
#         tot["GROSS WEIGHT"] = None
#         tot["NET WEIGHT"] = None
#         rows.append(tot)

#     return groups


# # ===========================================================================
# # UI
# # ===========================================================================
# json_path = st.sidebar.text_input("JSON file path", DEFAULT_JSON)

# if not Path(json_path).exists():
#     st.error(f"File not found: `{json_path}`")
#     st.stop()

# data = load_json(json_path)
# groups = build_bulk_pack_grouped(data)

# if not groups:
#     st.warning("No Bulk Pack tables found in the JSON.")
#     st.stop()


# all_frames: list[tuple[str, pd.DataFrame]] = []

# for color, rows in groups.items():
#     present_sizes = [s for s in SIZE_ORDER if any(s in r for r in rows)]
#     columns = LEFT_STATIC + present_sizes + RIGHT_STATIC
#     df = pd.DataFrame(rows).reindex(columns=columns)

#     st.subheader(f"🎨 {color}")
#     st.caption(f"{sum(1 for r in rows if r.get('CTN NO') != 'TOTAL')} size row(s)")

#     col_cfg = {
#         "CTN NO":          st.column_config.TextColumn("CTN NO", width="small"),
#         "CTN MES:":        st.column_config.NumberColumn("CTN MES:", width="small"),
#         "CTN QTY":         st.column_config.TextColumn("CTN QTY", width="small"),
#         "QTY":             st.column_config.NumberColumn("QTY", width="small"),
#         "COLOR":           st.column_config.TextColumn("COLOR", width="medium"),
#         "PREPACK STECKER": st.column_config.TextColumn("PREPACK STECKER", width="medium"),
#         "QTY PER CTN":     st.column_config.NumberColumn("QTY PER CTN", width="small"),
#         "TOTAL QTY":       st.column_config.TextColumn("TOTAL QTY", width="small"),
#         "GROSS WEIGHT":    st.column_config.NumberColumn(
#             "GROSS WEIGHT", width="small", format="%.2f"
#         ),
#         "NET WEIGHT":      st.column_config.NumberColumn(
#             "NET WEIGHT", width="small", format="%.2f"
#         ),
#     }
#     for s in present_sizes:
#         col_cfg[s] = st.column_config.NumberColumn(s, width="small")

#     try:
#         edited = st.data_editor(
#             data=df,
#             use_container_width=True,
#             hide_index=True,
#             num_rows="fixed",
#             column_config=col_cfg,
#             key=f"bulk_editor_{color}",
#         )
#     except TypeError:
#         st.warning(
#             "Your Streamlit version doesn't support data_editor with these options. "
#             "Showing read-only table. Upgrade with: pip install -U streamlit"
#         )
#         st.dataframe(df, use_container_width=True, hide_index=True)
#         edited = df

#     all_frames.append((color, edited))
#     st.divider()


# # ===========================================================================
# # Downloads
# # ===========================================================================
# st.markdown("### Download")

# export_rows: list[dict] = []
# for i, (color, frame) in enumerate(all_frames):
#     if i > 0:
#         export_rows.append({})
#     export_rows.extend(frame.to_dict(orient="records"))

# combined = pd.DataFrame(export_rows)

# col1, col2 = st.columns(2)

# with col1:
#     csv = combined.to_csv(index=False).encode("utf-8")
#     st.download_button(
#         "⬇ Download CSV (all colors)",
#         data=csv,
#         file_name="packing_list_bulk.csv",
#         mime="text/csv",
#         use_container_width=True,
#     )

# with col2:
#     try:
#         buf = io.BytesIO()
#         with pd.ExcelWriter(buf, engine="openpyxl") as writer:
#             for color, frame in all_frames:
#                 sheet = color[:31].replace("/", "-").replace("\\", "-")
#                 frame.to_excel(writer, index=False, sheet_name=sheet)
#         st.download_button(
#             "⬇ Download Excel (one sheet per color)",
#             data=buf.getvalue(),
#             file_name="packing_list_bulk.xlsx",
#             mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
#             use_container_width=True,
#         )
#     except ImportError:
#         st.info("Install `openpyxl` to enable Excel export:  pip install openpyxl")


# with st.sidebar:
#     st.markdown("### Colors detected")
#     for color, rows in groups.items():
#         n = sum(1 for r in rows if r.get("CTN NO") != "TOTAL")
#         st.write(f"- **{color}** — {n} size row(s)")




"""
bulk_pack.py
------------
Streamlit app: read parsed PO JSON → build BULK PACK packing list.
Each COLOR gets its own separate section, matching the yellow BULK PACK design.

Fixes:
    - TOTAL QTY column type mismatch (now always str)
    - Adds header block, section strip, PO header per color
    - Carton counter continues from Multi Pack + Single Pack

Run:
    streamlit run bulk_pack.py
"""

import io
import json
import math
from pathlib import Path

import pandas as pd
import streamlit as st

from multi_n_core import multi_n_carton_total

# ===========================================================================
# Config
# ===========================================================================
st.set_page_config(page_title="Packing List — Bulk Pack", layout="wide")
st.title("📦 Packing List — Bulk Pack")

DEFAULT_JSON = "61480360.json"
SIZE_ORDER = ["XS", "S", "M", "L", "XL", "XXL", "XXL+"]

DEFAULT_NW    = {"XS": 0.510, "S": 0.520, "M": 0.550, "L": 0.560, "XL": 0.610, "XXL": 0.650, "XXL+": 0.690}
DEFAULT_NNW   = {"XS": 0.490, "S": 0.500, "M": 0.530, "L": 0.540, "XL": 0.590, "XXL": 0.630, "XXL+": 0.670}
DEFAULT_EMPTY = {"XS": 0.49,  "S": 1.00,  "M": 1.00,  "L": 1.00,  "XL": 1.00,  "XXL": 1.00,  "XXL+": 1.00}

DEFAULT_CTN_MEAS = [
    "58.67 X 38.48 X 29.71 CM.G8. SL",
    "58.67 X 38.48 X 14.86 CM.G8S. SL",
    "38.48 X 29.33 X 14.86 CM.G-8M",
    "29.33 X 19.25 X 14.86 CM.G-8XS",
    "55.88 X 38.1  X 15.24 CM G13",
]

DEFAULT_BUYER = "OLD NAVY"
DEFAULT_STYLE = "905518-00-1"
DEFAULT_PO_NO = "61480360"

LEFT_STATIC  = ["CTN NO", "CTN MES:", "CTN QTY", "QTY", "COLOR", "PREPACK STECKER"]
RIGHT_STATIC = ["QTY PER CTN", "TOTAL QTY", "GROSS WEIGHT", "NET WEIGHT"]


# ===========================================================================
# Helpers
# ===========================================================================
@st.cache_data(show_spinner=False)
def load_json(path: str) -> dict:
    with open(path, "r", encoding="utf-8") as f:
        return json.load(f)


def color_display(raw: str) -> str:
    if not raw:
        return ""
    parts = raw.rsplit(" ", 1)
    return parts[-1] if len(parts) > 1 else raw


def color_code(raw: str) -> str:
    if not raw:
        return "000"
    first = raw.split()[0] if raw.split() else ""
    if "-" in first:
        return first.split("-")[-1]
    return "000"


def _int_or_none(v):
    if v is None:
        return None
    try:
        return int(float(str(v).strip().split()[0]))
    except (TypeError, ValueError, IndexError):
        return None


def _default_weights_df() -> pd.DataFrame:
    return pd.DataFrame(
        {
            "SIZE":          SIZE_ORDER,
            "N.W.":          [DEFAULT_NW[s]    for s in SIZE_ORDER],
            "N.N.W.":        [DEFAULT_NNW[s]   for s in SIZE_ORDER],
            "EMPTY CTN WT.": [DEFAULT_EMPTY[s] for s in SIZE_ORDER],
        }
    ).set_index("SIZE")


def _init_state():
    if "bulk_weights_df" not in st.session_state:
        st.session_state.bulk_weights_df = _default_weights_df()
    if "bulk_ctn_meas" not in st.session_state:
        st.session_state.bulk_ctn_meas = list(DEFAULT_CTN_MEAS)


# ===========================================================================
# Auto-detect preceding carton count (Multi Pack + Single Pack)
# ===========================================================================
def compute_preceding_ctn_total(data: dict) -> tuple[int, int]:
    """Returns (multi_total_ctn, single_total_ctn)."""
    multi = 0
    single = 0
    for _, table in data.get("tables", {}).items():
        rows = table.get("rows") or []
        if not rows:
            continue
        first = rows[0]
        ptype = first.get("PrePack Type")
        if ptype == "Multi" and first.get("Full Carton") == "Y":
            multi += _int_or_none((table.get("total") or {}).get("#PrePackets Ordered")) or 0
        elif ptype == "Single":
            single += _int_or_none((table.get("total") or {}).get("#PrePackets Ordered")) or 0
    # Multi Pack (N): several prepacks share a carton -> count real cartons
    multi += multi_n_carton_total(data)
    return multi, single


# ===========================================================================
# TOP BLOCK: Title + address + weight reference table
# ===========================================================================
def render_company_header(weights_df: pd.DataFrame) -> None:
    sizes = list(weights_df.index)
    nw_vals  = [f"{float(weights_df.loc[s,'N.W.']):.3f}"    for s in sizes]
    nnw_vals = [f"{float(weights_df.loc[s,'N.N.W.']):.3f}"  for s in sizes]

    size_cells = "".join(f"<th>{s}</th>" for s in sizes)
    nw_cells   = "".join(f"<td>{v}</td>" for v in nw_vals)
    nnw_cells  = "".join(f"<td>{v}</td>" for v in nnw_vals)

    html = f"""
    <style>
      .cc-title {{ text-align:center; font-family:'Courier New',monospace; font-weight:bold;
                   font-size:20px; text-decoration:underline; margin-bottom:2px; }}
      .cc-addr  {{ text-align:center; font-family:'Courier New',monospace; font-weight:bold;
                   font-size:15px; margin-bottom:12px; }}
      .cc-weight-wrap {{ display:flex; justify-content:center; margin-bottom:14px; }}
      .cc-weight {{ border-collapse:collapse; font-family:'Courier New',monospace; font-size:13px; }}
      .cc-weight th, .cc-weight td {{ border:1px solid #000; padding:2px 10px;
                                      text-align:center; min-width:48px; }}
      .cc-weight th.label {{ text-align:left; background:#fff; font-weight:bold; min-width:90px; }}
    </style>
    <div class="cc-title">CREATIVE COLLECTIONS LTD-1A.</div>
    <div class="cc-addr">Nishat Nagar , Tongi , Gazipur .</div>
    <div class="cc-weight-wrap">
      <table class="cc-weight">
        <tr><th class="label">SIZE</th><th class="label">SIZE</th>{size_cells}</tr>
        <tr><td class="label">N.W</td><td class="label">N.WT</td>{nw_cells}</tr>
        <tr><td class="label">N.N.W</td><td class="label">N.N.WT</td>{nnw_cells}</tr>
        <tr><td class="label">EMPTY CTN WET</td><td class="label"></td>
            <td colspan="{len(sizes)}"></td></tr>
      </table>
    </div>
    """
    st.markdown(html, unsafe_allow_html=True)


def render_section_strip(title: str) -> None:
    html = f"""
    <style>
      .cc-strip {{
        text-align:center; font-family:'Courier New',monospace; font-weight:bold;
        font-size:15px; letter-spacing:2px;
        background:#FFFF00; border:1px solid #000;
        padding:5px 0; margin:16px 0 10px 0;
      }}
    </style>
    <div class="cc-strip">{title}</div>
    """
    st.markdown(html, unsafe_allow_html=True)


def render_po_header(
    buyer: str, style: str, po_no: str,
    oiqty: int, ship_qty: int, ex_short: int, ctn_qty: int,
    ctn_meas_list: list, sku_item: str,
    unit_prepack: str, carton_label: str,
) -> None:
    meas_left = "".join(
        f'<div class="po-line"><span class="po-label">CTN MEAS.</span>'
        f'<span class="po-colon">:</span>'
        f'<span class="po-value">{m}</span></div>'
        for m in ctn_meas_list
    )
    html = f"""
    <style>
      .po-wrap {{ display:grid; grid-template-columns:1.4fr 1fr; gap:20px;
                 font-family:'Courier New',monospace; font-size:13.5px; line-height:1.5;
                 margin-bottom:10px; }}
      .po-line {{ display:flex; align-items:baseline; }}
      .po-label {{ display:inline-block; width:110px; font-weight:bold; }}
      .po-colon {{ width:14px; display:inline-block; text-align:center; }}
      .po-value {{ display:inline-block; }}
      .po-right {{ margin-top:60px; }}
    </style>
    <div class="po-wrap">
      <div class="po-left">
        <div class="po-line"><span class="po-label">BUYER</span><span class="po-colon">:</span><span class="po-value">{buyer}</span></div>
        <div class="po-line"><span class="po-label">STYLE</span><span class="po-colon">:</span><span class="po-value">{style}</span></div>
        <div class="po-line"><span class="po-label">P. O. #</span><span class="po-colon">:</span><span class="po-value">{po_no}</span></div>
        <div class="po-line"><span class="po-label">O/QTY</span><span class="po-colon">:</span><span class="po-value">{oiqty} PCS</span></div>
        <div class="po-line"><span class="po-label">SHIP QTY</span><span class="po-colon">:</span><span class="po-value">{ship_qty} PCS</span></div>
        <div class="po-line"><span class="po-label">EX/SHORT</span><span class="po-colon">:</span><span class="po-value">{ex_short} PCS</span></div>
        <div class="po-line"><span class="po-label">CTN QTY</span><span class="po-colon">:</span><span class="po-value">{ctn_qty} CTNS</span></div>
        {meas_left}
      </div>
      <div class="po-right">
        <div class="po-line"><span class="po-label">PO #</span><span class="po-colon">:</span><span class="po-value">{po_no}</span></div>
        <div class="po-line"><span class="po-label">SKU / Item</span><span class="po-colon">:</span><span class="po-value">{sku_item}</span></div>
        <div class="po-line"><span class="po-label">Unit / Prepack</span><span class="po-colon">:</span><span class="po-value">{unit_prepack}</span></div>
        <div class="po-line"><span class="po-label">CARTON</span><span class="po-colon">:</span><span class="po-value">{carton_label}</span></div>
      </div>
    </div>
    """
    st.markdown(html, unsafe_allow_html=True)


# ===========================================================================
# Build Bulk Pack grouped data
# ===========================================================================
def build_bulk_pack_grouped(
    data: dict, nw: dict, empty_weights: dict,
    multi_total: int, single_total: int, grand_total_ctn: int,
) -> tuple[dict, dict]:
    """
    Returns (groups, header_map):
      groups[color]      = list of row dicts (data + TOTAL)
      header_map[color]  = { oiqty, ship_qty, ex_short, ctn_qty, ... }
    """
    # ---- Group raw tables by color ----
    raw: dict[str, list] = {}
    for tname, table in data["tables"].items():
        trows = table.get("rows") or []
        if not trows:
            continue
        first = trows[0]
        if first.get("PrePack Type") != "Bulk":
            continue
        color = color_display(first.get("Universal CC #Color Desc", "")) or "UNKNOWN"
        raw.setdefault(color, []).append((tname, table))

    # ---- Compute header info ----
    header_map: dict[str, dict] = {}
    for color, tables in raw.items():
        oiqty_total = 0
        ship_qty_total = 0
        ctn_total = 0
        first_row = tables[0][1]["rows"][0]
        color_cd = color_code(first_row.get("Universal CC #Color Desc", ""))

        for _, table in tables:
            trows = table.get("rows") or []
            for r in trows:
                oiqty_total += _int_or_none(r.get("ORDER QTY")) or 0
                ship_qty_total += _int_or_none(r.get("SHIP QTY")) or 0
                pf = _int_or_none(r.get("Item Carton Pack Factor")) or 1
                ctn_total += math.ceil((_int_or_none(r.get("ORDER QTY")) or 0) / pf)

        header_map[color] = {
            "color_code": color_cd,
            "oiqty": oiqty_total,
            "ship_qty": ship_qty_total,
            "ex_short": max(oiqty_total - ship_qty_total, 0),
            "ctn_qty": ctn_total,
        }

    # ---- Build rows ----
    groups: dict[str, list] = {}
    global_cursor = multi_total + single_total + 1

    for color, tables in raw.items():
        rows: list[dict] = []
        tables_sorted = sorted(tables, key=lambda kv: kv[0])

        for _, table in tables_sorted:
            for r in (table.get("rows") or []):
                size = r.get("Size Desc")
                qty = _int_or_none(r.get("ORDER QTY")) or 0
                pf  = _int_or_none(r.get("Item Carton Pack Factor")) or 0
                if qty <= 0 or pf <= 0:
                    continue

                cartons = math.ceil(qty / pf)
                start = global_cursor
                end = start + cartons - 1
                global_cursor = end + 1

                nw_size = nw.get(size, 0.0)
                net_per_carton   = round(pf * nw_size, 2)
                empty_wt         = empty_weights.get(size, 0.0)
                gross_per_carton = round(net_per_carton + empty_wt, 2)

                row = {
                    "CTN NO":          start,          # Column A — start
                    "CTN MES:":        end,            # Column B — end
                    "CTN QTY":         str(cartons),   # str to keep column consistent
                    "QTY":             cartons,
                    "COLOR":           color if not rows else "",
                    "PREPACK STECKER": None,
                }
                for s in SIZE_ORDER:
                    row[s] = qty if s == size else None
                row["QTY PER CTN"]  = pf
                row["TOTAL QTY"]    = f"{qty} PCS"          # ALWAYS str
                row["GROSS WEIGHT"] = gross_per_carton
                row["NET WEIGHT"]   = net_per_carton
                rows.append(row)

        # ---- TOTAL row ----
        total_cartons = sum((r.get("QTY") or 0) for r in rows)
        tot_qty = sum(
            _int_or_none(str(r.get("TOTAL QTY", "")).replace("PCS", "").strip()) or 0
            for r in rows
        )
        tot = {
            "CTN NO":          "TOTAL",
            "CTN MES:":        None,
            "CTN QTY":         f"{total_cartons} CTN",
            "QTY":             None,
            "COLOR":           None,
            "PREPACK STECKER": None,
        }
        for s in SIZE_ORDER:
            pieces = sum((r.get(s) or 0) for r in rows)
            tot[s] = pieces if pieces else None
        tot["QTY PER CTN"]  = None
        tot["TOTAL QTY"]    = f"{tot_qty} PCS"          # ALWAYS str
        tot["GROSS WEIGHT"] = None
        tot["NET WEIGHT"]   = None
        rows.append(tot)

        groups[color] = rows

    return groups, header_map


# ===========================================================================
# Excel export — mirrors yellow BULK PACK layout
# ===========================================================================
def export_excel_bulk(
    groups: dict, header_map: dict, weights_df: pd.DataFrame,
    ctn_meas: list, buyer: str, style: str, po_no: str,
    grand_total_ctn: int,
) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "BULK PACK"

    thin = Side(style="thin", color="000000")
    med  = Side(style="medium", color="000000")
    border_all = Border(left=thin, right=thin, top=thin, bottom=thin)

    font_title   = Font(name="Courier New", size=18, bold=True, underline="single")
    font_addr    = Font(name="Courier New", size=13, bold=True)
    font_section = Font(name="Courier New", size=13, bold=True)
    font_bold    = Font(name="Courier New", size=11, bold=True)
    font_normal  = Font(name="Courier New", size=11)

    align_center = Alignment(horizontal="center", vertical="center")
    align_left   = Alignment(horizontal="left",   vertical="center")

    fill_header = PatternFill("solid", fgColor="D9D9D9")
    fill_yellow = PatternFill("solid", fgColor="FFFF00")
    fill_total  = PatternFill("solid", fgColor="F2F2F2")

    widths = {
        "A": 8,  "B": 8,  "C": 8,  "D": 8,  "E": 14, "F": 16,
        "G": 7,  "H": 7,  "I": 7,  "J": 7,  "K": 7,  "L": 7,  "M": 8,
        "N": 12, "O": 12, "P": 12, "Q": 12,
    }
    for col, w in widths.items():
        ws.column_dimensions[col].width = w

    # Title + address
    ws.merge_cells("A1:Q1")
    ws["A1"] = "CREATIVE COLLECTIONS LTD-1A."
    ws["A1"].font = font_title; ws["A1"].alignment = align_center

    ws.merge_cells("A2:Q2")
    ws["A2"] = "Nishat Nagar , Tongi , Gazipur ."
    ws["A2"].font = font_addr; ws["A2"].alignment = align_center

    # Weight reference
    ws["D4"] = "SIZE"; ws["D4"].font = font_bold; ws["D4"].border = border_all
    ws["F4"] = "SIZE"; ws["F4"].font = font_bold; ws["F4"].border = border_all
    for i, s in enumerate(SIZE_ORDER):
        c = ws.cell(row=4, column=7 + i, value=s)
        c.font = font_bold; c.alignment = align_center; c.border = border_all

    ws["D5"] = "N.W";  ws["D5"].font = font_bold; ws["D5"].border = border_all
    ws["F5"] = "N.WT"; ws["F5"].font = font_bold; ws["F5"].border = border_all
    for i, s in enumerate(SIZE_ORDER):
        c = ws.cell(row=5, column=7 + i, value=float(weights_df.loc[s, "N.W."]))
        c.number_format = "0.000"; c.alignment = align_center; c.border = border_all

    ws["D6"] = "N.N.W";  ws["D6"].font = font_bold; ws["D6"].border = border_all
    ws["F6"] = "N.N.WT"; ws["F6"].font = font_bold; ws["F6"].border = border_all
    for i, s in enumerate(SIZE_ORDER):
        c = ws.cell(row=6, column=7 + i, value=float(weights_df.loc[s, "N.N.W."]))
        c.number_format = "0.000"; c.alignment = align_center; c.border = border_all

    ws["D7"] = "EMPTY CTN WET"; ws["D7"].font = font_bold; ws["D7"].border = border_all

    r = 9
    for color, rows in groups.items():
        hdr = header_map.get(color, {})
        style_code = hdr.get("color_code", "000")
        style_full = f"{DEFAULT_STYLE.rsplit('-', 2)[0]}-{style_code}-1"

        # Yellow strip
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=17)
        c = ws.cell(row=r, column=1, value="BULK PACK")
        c.font = font_section; c.alignment = align_center
        c.fill = fill_yellow
        c.border = Border(top=med, bottom=med)
        r += 1

        # Left block
        po_left = [
            ("BUYER",      buyer),
            ("STYLE",      style_full),
            ("P. O. #",    po_no),
            ("O/QTY",      f"{hdr.get('oiqty', 0)} PCS"),
            ("SHIP QTY",   f"{hdr.get('ship_qty', 0)} PCS"),
            ("EX/SHORT",   f"{hdr.get('ex_short', 0)} PCS"),
            ("CTN QTY",    f"{hdr.get('ctn_qty', 0)} CTNS"),
        ]
        left_start = r
        for label, value in po_left:
            ws.cell(row=r, column=1, value=label).font = font_bold
            ws.cell(row=r, column=2, value=":").alignment = align_center
            ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=6)
            ws.cell(row=r, column=3, value=value).font = font_normal
            r += 1

        for m in ctn_meas[:4]:
            ws.cell(row=r, column=1, value="CTN MEAS.").font = font_bold
            ws.cell(row=r, column=2, value=":").alignment = align_center
            ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=6)
            ws.cell(row=r, column=3, value=m).font = font_normal
            r += 1

        # Right block
        right_start = left_start + 6
        first_ctn = rows[0].get("CTN NO", 1) if rows else 1
        po_right = [
            ("PO #",           po_no),
            ("SKU / Item",     style_full),
            ("Unit / Prepack", f"{hdr.get('oiqty', 0)}/{hdr.get('ctn_qty', 0)}"),
            ("CARTON",         f"{first_ctn} of {grand_total_ctn}"),
        ]
        rr = right_start
        for label, value in po_right:
            ws.cell(row=rr, column=12, value=label).font = font_bold
            ws.cell(row=rr, column=13, value=":").alignment = align_center
            ws.merge_cells(start_row=rr, start_column=14, end_row=rr, end_column=17)
            ws.cell(row=rr, column=14, value=value).font = font_normal
            rr += 1

        # Main table
        sizes = [s for s in SIZE_ORDER if any(s in row for row in rows)]
        size_cols = len(sizes)
        base_right = 7 + size_cols

        # Header row 1
        left_hdr = ["CTN", "CTN", "CTN", "QTY", "COLOR", "PREPACK STECKER"]
        for i, h in enumerate(left_hdr):
            c = ws.cell(row=r, column=1 + i, value=h)
            c.font = font_bold; c.alignment = align_center
            c.border = border_all; c.fill = fill_header

        ws.merge_cells(start_row=r, start_column=7, end_row=r, end_column=7 + size_cols - 1)
        c = ws.cell(row=r, column=7, value="SIZE")
        c.font = font_bold; c.alignment = align_center
        c.border = border_all; c.fill = fill_header

        right_hdr = ["QTY PER CTN", "TOTAL QTY", "GROSS WEIGHT", "NET WEIGHT"]
        for i, h in enumerate(right_hdr):
            c = ws.cell(row=r, column=base_right + i, value=h)
            c.font = font_bold; c.alignment = align_center
            c.border = border_all; c.fill = fill_header
        r += 1

        # Header row 2
        for i, s in enumerate(sizes):
            c = ws.cell(row=r, column=7 + i, value=s)
            c.font = font_bold; c.alignment = align_center
            c.border = border_all; c.fill = fill_header
        for i in range(6):
            ws.cell(row=r, column=1 + i).border = border_all
            ws.cell(row=r, column=1 + i).fill = fill_header
        for i in range(len(right_hdr)):
            ws.cell(row=r, column=base_right + i).border = border_all
            ws.cell(row=r, column=base_right + i).fill = fill_header
        r += 1

        # Data rows (exclude TOTAL)
        for row in rows[:-1]:
            ws.cell(row=r, column=1, value=row.get("CTN NO")).alignment = align_center
            ws.cell(row=r, column=2, value=row.get("CTN MES:")).alignment = align_center
            ws.cell(row=r, column=3, value=row.get("CTN QTY")).alignment = align_center
            ws.cell(row=r, column=4, value=row.get("QTY")).alignment = align_center
            ws.cell(row=r, column=5, value=row.get("COLOR", "")).alignment = align_center
            ws.cell(row=r, column=6, value=row.get("PREPACK STECKER", "")).alignment = align_center
            for i, s in enumerate(sizes):
                v = row.get(s)
                if v is not None:
                    ws.cell(row=r, column=7 + i, value=v).alignment = align_center
            ws.cell(row=r, column=base_right + 0, value=row.get("QTY PER CTN")).alignment = align_center
            ws.cell(row=r, column=base_right + 1, value=row.get("TOTAL QTY")).alignment = align_center
            ws.cell(row=r, column=base_right + 2, value=row.get("GROSS WEIGHT")).alignment = align_center
            ws.cell(row=r, column=base_right + 2).number_format = "0.00"
            ws.cell(row=r, column=base_right + 3, value=row.get("NET WEIGHT")).alignment = align_center
            ws.cell(row=r, column=base_right + 3).number_format = "0.00"
            for c_idx in range(1, base_right + 4):
                ws.cell(row=r, column=c_idx).border = border_all
            r += 1

        # TOTAL row
        tot = rows[-1]
        ws.cell(row=r, column=1, value="TOTAL").font = font_bold
        ws.cell(row=r, column=1).alignment = align_center
        ws.cell(row=r, column=1).fill = fill_total
        ws.cell(row=r, column=3, value=tot.get("CTN QTY")).alignment = align_center
        ws.cell(row=r, column=3).fill = fill_total
        for i, s in enumerate(sizes):
            v = tot.get(s)
            if v is not None:
                c = ws.cell(row=r, column=7 + i, value=v)
                c.alignment = align_center; c.fill = fill_total
        ws.cell(row=r, column=base_right + 1, value=tot.get("TOTAL QTY")).alignment = align_center
        ws.cell(row=r, column=base_right + 1).fill = fill_total
        for c_idx in range(1, base_right + 4):
            ws.cell(row=r, column=c_idx).border = border_all
        r += 2

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ===========================================================================
# UI
# ===========================================================================
_init_state()

with st.sidebar:
    st.markdown("### ⚖️ Weight settings (kg)")
    edited_w = st.data_editor(
        st.session_state.bulk_weights_df,
        use_container_width=True, num_rows="fixed",
        column_config={
            "N.W.":          st.column_config.NumberColumn("N.W.",          format="%.3f", step=0.001),
            "N.N.W.":        st.column_config.NumberColumn("N.N.W.",        format="%.3f", step=0.001),
            "EMPTY CTN WT.": st.column_config.NumberColumn("EMPTY CTN WT.", format="%.3f", step=0.001),
        },
        key="bulk_weights_editor",
    )
    if not edited_w.equals(st.session_state.bulk_weights_df):
        st.session_state.bulk_weights_df = edited_w
        st.rerun()

    if st.button("↺ Reset weights", use_container_width=True):
        st.session_state.bulk_weights_df = _default_weights_df()
        st.rerun()

    st.divider()
    st.markdown("### 📦 CTN MEAS.")
    for i, m in enumerate(st.session_state.bulk_ctn_meas):
        new_m = st.text_input(f"Line {i+1}", value=m, key=f"bulk_ctn_meas_{i}")
        if new_m != m:
            st.session_state.bulk_ctn_meas[i] = new_m

    st.divider()
    json_path = st.text_input("JSON file path", DEFAULT_JSON)


if not Path(json_path).exists():
    st.error(f"File not found: `{json_path}`")
    st.stop()

data = load_json(json_path)
multi_total, single_total = compute_preceding_ctn_total(data)
GRAND_TOTAL_CTN = multi_total + single_total + sum(
    math.ceil((_int_or_none(r.get("ORDER QTY")) or 0) / (_int_or_none(r.get("Item Carton Pack Factor")) or 1))
    for t in data.get("tables", {}).values()
    for r in (t.get("rows") or [])
    if (t.get("rows") or [{}])[0].get("PrePack Type") == "Bulk"
)

weights = {
    "NW":    {s: float(st.session_state.bulk_weights_df.loc[s, "N.W."])          for s in st.session_state.bulk_weights_df.index},
    "EMPTY": {s: float(st.session_state.bulk_weights_df.loc[s, "EMPTY CTN WT."]) for s in st.session_state.bulk_weights_df.index},
}

groups, header_map = build_bulk_pack_grouped(
    data, weights["NW"], weights["EMPTY"], multi_total, single_total, GRAND_TOTAL_CTN,
)

if not groups:
    st.warning("No Bulk Pack tables found in the JSON.")
    st.stop()

po_header = data.get("PO_HEADER", {}) or {}
buyer_fixed = po_header.get("BUYER") or DEFAULT_BUYER
po_no_fixed = po_header.get("P.O._#") or DEFAULT_PO_NO

# ===========================================================================
# Render
# ===========================================================================
render_company_header(st.session_state.bulk_weights_df)

all_frames: list = []
for color, rows in groups.items():
    sizes = [s for s in SIZE_ORDER if any(s in r for r in rows)]
    columns = LEFT_STATIC + sizes + RIGHT_STATIC
    df = pd.DataFrame(rows).reindex(columns=columns)

    hdr = header_map.get(color, {})
    style_full = f"{DEFAULT_STYLE.rsplit('-', 2)[0]}-{hdr.get('color_code', '000')}-1"
    unit_prepack = f"{hdr.get('oiqty', 0)}/{hdr.get('ctn_qty', 0)}"
    first_ctn = rows[0].get("CTN NO", 1) if rows else 1
    carton_label = f"{first_ctn} of {GRAND_TOTAL_CTN}"

    render_section_strip("BULK PACK")
    render_po_header(
        buyer=buyer_fixed, style=style_full, po_no=po_no_fixed,
        oiqty=hdr.get("oiqty", 0), ship_qty=hdr.get("ship_qty", 0),
        ex_short=hdr.get("ex_short", 0), ctn_qty=hdr.get("ctn_qty", 0),
        ctn_meas_list=st.session_state.bulk_ctn_meas[:4],
        sku_item=style_full, unit_prepack=unit_prepack, carton_label=carton_label,
    )

    col_cfg = {
        "CTN NO":          st.column_config.NumberColumn("CTN NO", width="small", disabled=True),
        "CTN MES:":        st.column_config.NumberColumn("CTN MES:", width="small", disabled=True),
        "CTN QTY":         st.column_config.TextColumn("CTN QTY", width="small", disabled=True),
        "QTY":             st.column_config.NumberColumn("QTY", width="small", disabled=True),
        "COLOR":           st.column_config.TextColumn("COLOR", width="medium", disabled=True),
        "PREPACK STECKER": st.column_config.TextColumn("PREPACK STECKER", width="medium", disabled=True),
        "QTY PER CTN":     st.column_config.NumberColumn("QTY PER CTN", width="small", disabled=True),
        # ↓ FIX: use TextColumn consistently (all values are strings)
        "TOTAL QTY":       st.column_config.TextColumn("TOTAL QTY", width="small", disabled=True),
        "GROSS WEIGHT":    st.column_config.NumberColumn("GROSS WEIGHT", width="small", format="%.2f", disabled=True),
        "NET WEIGHT":      st.column_config.NumberColumn("NET WEIGHT",   width="small", format="%.2f", disabled=True),
    }
    for s in sizes:
        col_cfg[s] = st.column_config.NumberColumn(s, width="small", disabled=True)

    try:
        edited = st.data_editor(
            data=df, use_container_width=True, hide_index=True,
            num_rows="fixed", column_config=col_cfg,
            key=f"bulk_editor_{color}",
        )
    except TypeError:
        st.dataframe(df, use_container_width=True, hide_index=True)
        edited = df

    all_frames.append((color, edited))
    st.divider()


# ===========================================================================
# Downloads
# ===========================================================================
st.markdown("### Download")

export_rows: list[dict] = []
for i, (color, frame) in enumerate(all_frames):
    if i > 0:
        export_rows.append({})
    export_rows.extend(frame.to_dict(orient="records"))
combined = pd.DataFrame(export_rows)

col1, col2 = st.columns(2)
with col1:
    st.download_button(
        "⬇ Download CSV (all colors)",
        data=combined.to_csv(index=False).encode("utf-8"),
        file_name="packing_list_bulk.csv",
        mime="text/csv", use_container_width=True,
    )
with col2:
    try:
        xlsx = export_excel_bulk(
            groups=groups, header_map=header_map,
            weights_df=st.session_state.bulk_weights_df,
            ctn_meas=st.session_state.bulk_ctn_meas,
            buyer=buyer_fixed, style=DEFAULT_STYLE, po_no=po_no_fixed,
            grand_total_ctn=GRAND_TOTAL_CTN,
        )
        st.download_button(
            "⬇ Download Excel (BULK PACK format)",
            data=xlsx, file_name="packing_list_bulk.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            use_container_width=True,
        )
    except Exception as e:
        st.error(f"Excel export failed: {e}")


with st.sidebar:
    st.markdown("### Colors detected")
    for color, rows in groups.items():
        n = sum(1 for r in rows if r.get("CTN NO") != "TOTAL")
        st.write(f"- **{color}** — {n} size row(s)")