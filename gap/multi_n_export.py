"""
multi_n_export.py
-----------------
Excel export (sheets 'PKL' + 'sum') for MULTI PACK ( N ).
Mirrors:  Style - 908546 PO - 61429740 - Qty - 6431 Pcs  Multi.xlsx

Layout (same columns as the reference, 7 size columns G..M):
    A CTN NO (from)   B CTN NO (to)   C CTN MES   D CTN QTY   E COLOR
    F PREPACK STECKER G..M sizes (units per prepack)
    N PER BLST (=sum of sizes)  O 'X'  P prepacks per carton
    Q QTY PER CTN (=N*P)  R TOTAL QTY (=Q*D)  S GROSS WEIGHT  T NET WEIGHT

All derived cells are live Excel formulas, so editing CTN QTY / the size units /
the weight table recalculates the sheet.
"""

import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L

from multi_n_core import size_family

FONT = "Courier New"


def export_multi_n_excel(
    sections: list,
    nw: dict,
    nnw: dict,
    tare_by_code: dict,
    net_deduct: float,
    ctn_meas: list,
    buyer: str,
    po_no: str,
    grand_total_ctn: int,
    season: str = "Sum'26",
) -> bytes:
    thin = Side(style="thin", color="000000")
    med = Side(style="medium", color="000000")
    box = Border(left=thin, right=thin, top=thin, bottom=thin)

    f_title = Font(name=FONT, size=18, bold=True, underline="single")
    f_addr = Font(name=FONT, size=13, bold=True)
    f_bold = Font(name=FONT, size=11, bold=True)
    f_norm = Font(name=FONT, size=11)
    ctr = Alignment(horizontal="center", vertical="center")
    left = Alignment(horizontal="left", vertical="center")
    fill_h = PatternFill("solid", fgColor="D9D9D9")
    fill_t = PatternFill("solid", fgColor="F2F2F2")

    all_sizes = [s for sec in sections for s in sec["sizes"]]
    cols = size_family(all_sizes)[:7]          # sheet size columns (G..M)
    SC0 = 7                                    # first size column (G)
    SC1 = SC0 + 6                              # last size column  (M)
    cPB, cX, cP, cQ, cR, cS, cT = 14, 15, 16, 17, 18, 19, 20   # N..T

    wb = Workbook()
    ws = wb.active
    ws.title = "PKL"

    for col, w in {"A": 7, "B": 7, "C": 8, "D": 10, "E": 20, "F": 18,
                   "G": 8, "H": 8, "I": 8, "J": 8, "K": 8, "L": 8, "M": 8,
                   "N": 12, "O": 4, "P": 8, "Q": 12, "R": 12, "S": 14, "T": 14}.items():
        ws.column_dimensions[col].width = w

    # ---- title ----------------------------------------------------------
    ws.merge_cells("A1:T1")
    ws["A1"] = "CREATIVE COLLECTIONS LTD-1A."
    ws["A1"].font, ws["A1"].alignment = f_title, ctr
    ws.merge_cells("A2:T2")
    ws["A2"] = "Nishat Nagar , Tongi , Gazipur ."
    ws["A2"].font, ws["A2"].alignment = f_addr, ctr

    # ---- weight reference table (rows 4-7) ------------------------------
    for r, (a, b) in zip((4, 5, 6, 7), (("SIZE", "SIZE"), ("N.W", "N.WT"),
                                         ("N.N.W", "N.N.WT"), ("EMPTY CTN WET", ""))):
        ws.cell(r, 4, a).font = f_bold
        ws.cell(r, 6, b).font = f_bold
        for c in range(4, SC1 + 1):
            ws.cell(r, c).border = box
    for i, s in enumerate(cols):
        c = SC0 + i
        ws.cell(4, c, s).font = f_bold
        ws.cell(4, c).alignment = ctr
        ws.cell(5, c, nw.get(s)).number_format = "0.000"
        ws.cell(6, c, nnw.get(s)).number_format = "0.000"
        ws.cell(5, c).alignment = ws.cell(6, c).alignment = ctr
    ws.merge_cells("D7:E7")

    # carton tare table  N4:T7  (same place as dummy_61480360_packing_list.xlsx)
    codes = list(tare_by_code)[:4]
    for rr_, lab_ in ((4, "CTN CODE"), (5, "EMPTY CTN WT"), (6, "GROSS-NET MULTI")):
        ws.merge_cells(start_row=rr_, start_column=14, end_row=rr_, end_column=16)
        ws.cell(rr_, 14, lab_).font = f_bold
    for i, code in enumerate(codes):
        ws.cell(4, 17 + i, code).font = f_bold
        ws.cell(5, 17 + i, tare_by_code[code]).number_format = "0.00"
        ws.cell(4, 17 + i).alignment = ws.cell(5, 17 + i).alignment = ctr
    ws.cell(6, 17, net_deduct).number_format = "0.00"   # $Q$6
    for r in (4, 5, 6):
        for c in range(14, 21):
            ws.cell(r, c).border = box
    TARE_CODES, TARE_VALS, NET_CELL = "$Q$4:$T$4", "$Q$5:$T$5", "$Q$6"

    # ---- sections -------------------------------------------------------
    r = 9
    prev_last_row = None
    sum_refs = []                      # for 'sum' sheet
    for sec in sections:
        # strip
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=20)
        c = ws.cell(r, 1, "MULTY PACK ( N )")
        c.font, c.alignment = Font(name=FONT, size=13, bold=True), ctr
        c.border = Border(top=med, bottom=med)
        strip = r
        r += 1

        hdr0 = r
        po_left = ["BUYER", "STYLE", "P. O. #", "O/QTY", "SHIP QTY", "EX/SHORT", "CTN QTY"]
        for k, lab in enumerate(po_left):
            ws.merge_cells(start_row=r + k, start_column=1, end_row=r + k, end_column=2)
            ws.cell(r + k, 1, lab).font = f_bold
            ws.cell(r + k, 3, ":").alignment = ctr
            ws.merge_cells(start_row=r + k, start_column=4, end_row=r + k, end_column=8)
            ws.cell(r + k, 4).alignment = left
            ws.cell(r + k, 4).font = f_norm
        r += len(po_left)
        for m in ctn_meas:
            ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)
            ws.cell(r, 1, "CTN MEAS.").font = f_bold
            ws.cell(r, 3, ":").alignment = ctr
            ws.merge_cells(start_row=r, start_column=4, end_row=r, end_column=8)
            ws.cell(r, 4, m).font = f_norm
            r += 1

        # table header
        th = r
        for i, h in enumerate(["CTN", "", "CTN", "CTN QTY", "COLOR", "PREPACK STECKER"]):
            ws.cell(th, 1 + i, h)
        ws.merge_cells(start_row=th, start_column=SC0, end_row=th, end_column=SC1)
        ws.cell(th, SC0, "SIZE")
        for c, h in ((cPB, "PER BLST"), (cQ, "QTY PER CTN"), (cR, "TOTAL QTY"),
                     (cS, "GROSS WEIGHT"), (cT, "NET WEIGHT")):
            ws.cell(th, c, h)
        ws.merge_cells(start_row=th, start_column=cPB, end_row=th + 1, end_column=cP)
        ws.cell(th + 1, 1, "NO")
        ws.cell(th + 1, 3, "MES:")
        ws.cell(th + 1, 5, "COLOR")
        for i in range(7):
            ws.cell(th + 1, SC0 + i, f"={L(SC0 + i)}$4")
        for rr in (th, th + 1):
            for c in range(1, cT + 1):
                cell = ws.cell(rr, c)
                cell.font, cell.fill, cell.border = f_bold, fill_h, box
                cell.alignment = ctr
        r = th + 2

        # data rows
        first_data = r
        for k, row in enumerate(sec["rows"]):
            # CTN NO from / to
            if k == 0 and prev_last_row is None:
                ws.cell(r, 1, row["ctn_from"])
            elif k == 0:
                ws.cell(r, 1, f"=B{prev_last_row}+1")
            else:
                ws.cell(r, 1, f"=B{r - 1}+1")
            ws.cell(r, 2, f"=A{r}+D{r}-1")
            ws.cell(r, 3, row["ctn_mes"])
            ws.cell(r, 4, row["ctn_qty"])
            if k == 0:
                ws.cell(r, 5, sec["color"])
                ws.cell(r, 6, int(sec["prepack"]))
            for i, s in enumerate(cols):
                u = sec["units"].get(s)
                if u:
                    ws.cell(r, SC0 + i, u)
            ws.cell(r, cPB, f"=SUM({L(SC0)}{r}:{L(SC1)}{r})")
            ws.cell(r, cX, "X")
            ws.cell(r, cP, row["prepacks_per_ctn"])
            ws.cell(r, cQ, f"={L(cPB)}{r}*{L(cP)}{r}")
            ws.cell(r, cR, f"={L(cQ)}{r}*D{r}")
            ws.cell(r, cS,
                    f"=SUMPRODUCT({L(SC0)}{r}:{L(SC1)}{r},$G$5:$M$5)*{L(cP)}{r}"
                    f"+IFERROR(INDEX({TARE_VALS},MATCH(C{r},{TARE_CODES},0)),0)")
            ws.cell(r, cT, f"={L(cS)}{r}-{NET_CELL}")
            ws.cell(r, cS).number_format = ws.cell(r, cT).number_format = "0.00"
            for c in range(1, cT + 1):
                ws.cell(r, c).border = box
                ws.cell(r, c).alignment = ctr
                ws.cell(r, c).font = f_norm
            r += 1
        last_data = r - 1
        prev_last_row = last_data

        # TOTAL row
        tr = r
        ws.cell(tr, 1, "TOTAL")
        ws.cell(tr, 4, f"=SUM(D{first_data}:D{last_data})")
        ws.cell(tr, 5, "CTNS")
        for i, s in enumerate(cols):
            if sec["units"].get(s):
                col = L(SC0 + i)
                ws.cell(tr, SC0 + i,
                        f"=SUMPRODUCT({col}{first_data}:{col}{last_data},"
                        f"{L(cP)}{first_data}:{L(cP)}{last_data},"
                        f"D{first_data}:D{last_data})")
        ws.cell(tr, cR, f"=SUM({L(cR)}{first_data}:{L(cR)}{last_data})")
        ws.cell(tr, cS, "PCS")
        for c in range(1, cT + 1):
            cell = ws.cell(tr, c)
            cell.font, cell.fill, cell.border, cell.alignment = f_bold, fill_t, box, ctr

        # header values (need total-row refs)
        vals = [
            buyer,
            sec["style_full"],
            po_no,
            f"{sec['qty_ordered']}",
            f"=R{tr}",
            f"=D{hdr0 + 4}-D{hdr0 + 3}",
            f"=D{tr}",
        ]
        # put O/QTY as number + unit text via number format
        for k, v in enumerate(vals):
            cell = ws.cell(hdr0 + k, 4)
            cell.value = int(v) if k == 3 else v
            if k >= 3:
                cell.number_format = '#,##0 "PCS"' if k < 6 else '#,##0 "CTNS"'
            cell.font, cell.alignment = f_norm, left
        # right block (PO # .. CARTON) beside O/QTY..CTN QTY lines
        carton_label = (f'=TEXT(A{first_data},"00")&" of {grand_total_ctn}"')
        right = [("PO #", f"=D{hdr0 + 2}"),
                 ("SKU / Item", f"=D{hdr0 + 1}"),
                 ("Unit / Prepack", f"{sec['units_per_prepack']}/{sec['per_ctn']}"),
                 ("CARTON", carton_label)]
        for k, (lab, val) in enumerate(right):
            rr = hdr0 + 3 + k
            ws.merge_cells(start_row=rr, start_column=11, end_row=rr, end_column=12)
            ws.cell(rr, 11, lab).font = f_bold
            ws.cell(rr, 13, ":").alignment = ctr
            ws.merge_cells(start_row=rr, start_column=14, end_row=rr, end_column=17)
            ws.cell(rr, 14, val).font = f_norm
            ws.cell(rr, 14).alignment = left

        sum_refs.append({
            "color": sec["color"], "qty": sec["qty_ordered"],
            "ship": f"=PKL!R{tr}", "ctns": f"=PKL!D{tr}",
        })
        r = tr + 2

    # ---- sheet 'sum' ----------------------------------------------------
    s2 = wb.create_sheet("sum")
    for col, w in zip("ABCDEFGHIJ", (8, 14, 10, 20, 12, 12, 10, 10, 10, 10)):
        s2.column_dimensions[col].width = w
    s2.merge_cells("A3:J3")
    s2["A3"] = "CREATIVE COLLECTIONS LIMITED (UNIT-1A)"
    s2["A3"].font, s2["A3"].alignment = f_addr, ctr
    s2.merge_cells("A4:J4")
    s2["A4"] = "Nishatnagar,Tongi,Gazipur"
    s2["A4"].alignment = ctr
    s2.merge_cells("A5:J5")
    s2["A5"] = "PACKING LIST SUMMARY"
    s2["A5"].font, s2["A5"].alignment = f_bold, ctr
    s2["A7"], s2["B7"] = "BUYER", buyer
    s2["A8"], s2["B8"] = "STYLE", sections[0]["style_no"] if sections else ""
    s2["A7"].font = s2["A8"].font = f_bold
    for i, h in enumerate(["DC", "PO NO", "SEASON", "COLOR/CODE", "ORD/QTY", "SHIP/QTY",
                           "EXCES", "SHORT", "PER (%)", "CTNS"]):
        c = s2.cell(10, 1 + i, h)
        c.font, c.fill, c.border, c.alignment = f_bold, fill_h, box, ctr
    rr = 11
    for ref in sum_refs:
        s2.cell(rr, 1, "BDC")
        s2.cell(rr, 2, po_no)
        s2.cell(rr, 3, season)
        s2.cell(rr, 4, ref["color"])
        s2.cell(rr, 5, ref["qty"])
        s2.cell(rr, 6, ref["ship"])
        s2.cell(rr, 7, f"=MAX(F{rr}-E{rr},0)")
        s2.cell(rr, 8, f"=MAX(E{rr}-F{rr},0)")
        s2.cell(rr, 9, f"=IF(E{rr}=0,0,(F{rr}-E{rr})/E{rr})")
        s2.cell(rr, 9).number_format = "0.00%"
        s2.cell(rr, 10, ref["ctns"])
        for c in range(1, 11):
            s2.cell(rr, c).border, s2.cell(rr, c).alignment = box, ctr
        rr += 1
    s2.cell(rr, 4, "TOTAL")
    for c, col in ((5, "E"), (6, "F"), (7, "G"), (8, "H"), (10, "J")):
        s2.cell(rr, c, f"=SUM({col}11:{col}{rr - 1})")
    s2.cell(rr, 9, f"=IF(E{rr}=0,0,(F{rr}-E{rr})/E{rr})")
    s2.cell(rr, 9).number_format = "0.00%"
    for c in range(1, 11):
        cell = s2.cell(rr, c)
        cell.font, cell.fill, cell.border, cell.alignment = f_bold, fill_t, box, ctr

    # print setup: landscape, one page wide
    from openpyxl.worksheet.properties import PageSetupProperties
    for sh in (ws, s2):
        sh.page_setup.orientation = "landscape"
        sh.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
        sh.page_setup.fitToWidth = 1
        sh.page_setup.fitToHeight = 0

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
